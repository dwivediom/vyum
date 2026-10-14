"""Spreadsheet edit eval: the task types SpreadsheetBench / SheetCopilot found hard, on a realistic workbook.
Each task runs on a fresh copy; the result is read back (values IronCalc computed) and checked against Python.

    .venv/bin/python -m evals.sheet_tasks [--only N] [--out evals/results_sheets.json]
"""
import argparse
import datetime as dt
import json
import random
import shutil
import tempfile
import time
from pathlib import Path

import openpyxl
from python_calamine import CalamineWorkbook

from engine import sheets
from engine.office import OfficeError

REPS = {"Asha Verma": "Sales", "Rahul Mehta": "Sales", "Neha Iyer": "Partnerships", "Vikram Rao": "Enterprise",
        "Sara Khan": "Enterprise"}
REGIONS = ["North", "South", "East", "West"]


def build(path: Path):
    rnd = random.Random(7)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Date", "Region", "Rep", "Units", "Unit Price", "Revenue", "Revenue with GST"])
    start = dt.date(2025, 1, 3)
    for i in range(40):
        r = i + 2
        d = start + dt.timedelta(days=rnd.randint(0, 120))
        price = rnd.choice([120, 180, 250, 400, None]) if i % 9 else None
        ws.append([d, rnd.choice(REGIONS), rnd.choice(list(REPS)), rnd.randint(5, 160), price, f"=D{r}*E{r}",
                   f"=F{r}*(1+#REF!)"])
        ws.cell(r, 1).number_format = "yyyy-mm-dd"
    ws.append(["Total", None, None, "=SUM(D2:D41)", None, "=SUM(F2:F41)", None])
    st = wb.create_sheet("Staff")
    st.append(["Rep", "Department", "Joined"])
    for name, dep in REPS.items():
        st.append([name, dep, dt.date(2020 + len(name) % 5, 4, 1)])
    bd = wb.create_sheet("Bands")
    bd.append(["Min Units", "Band"])
    for lo, name in [(0, "Small"), (50, "Medium"), (100, "Large"), (150, "XL")]:
        bd.append([lo, name])
    yr = wb.create_sheet("Years")
    yr.append(["Year", "Revenue"])
    for y, v in zip(range(2019, 2025), [1_200_000, 1_380_000, 1_510_000, 1_790_000, 2_050_000, 2_400_000]):
        yr.append([y, v])
    wb.save(path)


def values(path: Path, sheet="Sales"):
    return [[None if v == "" else v for v in r] for r in sheets.read_values(path, sheet)]


def col(rows, header):
    j = rows[0].index(header)
    return j


def find_col(rows, *words):
    for j, h in enumerate(rows[0]):
        if h and all(w.lower() in str(h).lower() for w in words):
            return j
    return None


def data(rows):
    return [r for r in rows[1:] if r and r[0] != "Total" and any(v is not None for v in r[:6])]


def near(a, b, tol=1e-6):
    try:
        return abs(float(a) - float(b)) <= tol * max(1.0, abs(float(b)))
    except (TypeError, ValueError):
        return False


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
    return None


def cell(rows, addr):
    import re
    m = re.fullmatch(r"([A-Z]+)(\d+)", addr)
    c = sheets.col_num(m.group(1)) - 1
    r = int(m.group(2)) - 1
    return rows[r][c] if r < len(rows) and c < len(rows[r]) else None


# each check gets (rows before, rows after, path) and returns (ok, note)
def t_share(b, a, p):
    j = find_col(a, "share")
    if j is None:
        return False, "no share column"
    d = data(a)
    rev = [r[5] for r in d]
    total = sum(x for x in rev if isinstance(x, (int, float)))
    ok = all(near(r[j], (r[5] or 0) / total) for r in d)
    return ok, f"sum of shares {sum((r[j] or 0) for r in d):.4f}"


def t_sumifs(b, a, p):
    want = sum((r[3] or 0) * (r[4] or 0) for r in data(b) if r[1] == "North" and as_date(r[0]).month == 2)
    got = cell(a, "J2")
    return near(got, want), f"J2={got} want {want}"


def t_lookup(b, a, p):
    j = find_col(a, "department")
    if j is None:
        return False, "no Department column"
    bad = [r for r in data(a) if r[j] != REPS[r[2]]]
    return not bad, f"{len(bad)} wrong"


def t_running(b, a, p):
    j = find_col(a, "running") or find_col(a, "cumulative")
    if j is None:
        return False, "no running total column"
    acc, ok = 0, True
    for r in data(a):
        acc += r[5] or 0
        ok &= near(r[j], acc)
    return ok, f"last {data(a)[-1][j]} want {acc}"


def t_rank(b, a, p):
    j = find_col(a, "rank")
    if j is None:
        return False, "no Rank column"
    d = data(a)
    rev = sorted((r[5] or 0 for r in d), reverse=True)
    ok = all(int(r[j]) == rev.index(r[5] or 0) + 1 for r in d if isinstance(r[j], (int, float)))
    return ok and all(isinstance(r[j], (int, float)) for r in d), "ranks checked"


def t_band(b, a, p):
    j = find_col(a, "size")
    if j is None:
        return False, "no Size column"
    want = lambda u: "Large" if u >= 100 else "Medium" if u >= 50 else "Small"
    bad = [r for r in data(a) if str(r[j]).strip() != want(r[3])]
    return not bad, f"{len(bad)} wrong"


def t_pivot(b, a, p):
    want = {g: sum((r[3] or 0) * (r[4] or 0) for r in data(b) if r[1] == g) for g in REGIONS}
    cnt = {g: sum(1 for r in data(b) if r[1] == g) for g in REGIONS}
    found = {}
    for n in sheets.sheet_names(p):
        for r in values(p, n):
            for j, v in enumerate(r[7 if n == "Sales" else 0:], 7 if n == "Sales" else 0):
                if v in REGIONS and j + 2 < len(r) and isinstance(r[j + 1], (int, float)):
                    found[v] = (r[j + 1], r[j + 2])
    ok = len(found) == 4 and all(
        (near(found[g][0], want[g]) and near(found[g][1], cnt[g])) or (near(found[g][1], want[g]) and near(found[g][0], cnt[g]))
        for g in REGIONS)
    return ok, f"found {found}"


def t_due(b, a, p):
    j = find_col(a, "due")
    if j is None:
        return False, "no Due column"
    import calendar
    bad = 0
    for r in data(a):
        d = as_date(r[0])
        y, m = (d.year + (d.month // 12), d.month % 12 + 1)
        want = dt.date(y, m, calendar.monthrange(y, m)[1])
        bad += as_date(r[j]) != want
    return bad == 0, f"{bad} wrong"


def t_first(b, a, p):
    j = find_col(a, "first")
    if j is None:
        return False, "no First Name column"
    bad = [r for r in data(a) if r[j] != r[2].split()[0]]
    return not bad, f"{len(bad)} wrong"


def t_pmt(b, a, p):
    r, n, pv = 0.09 / 12, 60, 500000
    want = pv * r / (1 - (1 + r) ** -n)
    got = cell(a, "J5")
    return near(abs(got or 0), want, 1e-4), f"J5={got} want {want:.2f}"


def t_avg_blank(b, a, p):
    prices = [r[4] for r in data(b) if isinstance(r[4], (int, float))]
    got = cell(a, "J8")
    return near(got, sum(prices) / len(prices)), f"J8={got} want {sum(prices) / len(prices):.4f}"


def t_fix(b, a, p):
    d = data(a)
    ok = all(near(r[6], (r[5] or 0) * 1.18) for r in d)
    return ok, f"G2={d[0][6]} want {(d[0][5] or 0) * 1.18}"


def t_add_row(b, a, p):
    d = data(a)
    new = [r for r in d if r[2] == "Asha Verma" and r[3] == 40 and r[4] == 250]
    tot = next(r for r in a if r and r[0] == "Total")
    ok = bool(new) and near(new[0][5], 10000) and near(tot[5], sum(r[5] or 0 for r in d)) and near(tot[3], sum(r[3] or 0 for r in d))
    return ok, f"new={bool(new)} total F={tot[5]}"


def t_sort(b, a, p):
    rev = [r[5] or 0 for r in data(a)]
    tot = next(r for r in a if r and r[0] == "Total")
    return rev == sorted(rev, reverse=True) and near(tot[5], sum(rev)), f"first {rev[:3]}"


def t_cagr(b, a, p):
    got = cell(values(p, "Years"), "D2")
    want = (2_400_000 / 1_200_000) ** (1 / 5) - 1
    return near(got, want, 1e-4) or near(got, want * 100, 1e-4), f"D2={got} want {want:.6f}"


DEPT_RATE = {"Sales": 0.05, "Enterprise": 0.08}


def rev(r):
    return (r[3] or 0) * (r[4] or 0)


def pairs(a, start_col=7):
    """Every label -> number pair to the right of the data, wherever the model put its summary table."""
    out = {}
    for r in a:
        for j in range(start_col, len(r) - 1):
            if isinstance(r[j], str) and isinstance(r[j + 1], (int, float)):
                out.setdefault(r[j], []).append(r[j + 1])
    return out


def all_pairs(p):
    """Label -> numbers anywhere in the workbook (a summary may be put on a sheet of its own)."""
    out = {}
    for n in sheets.sheet_names(p):
        for k, v in pairs(values(p, n), 0 if n != "Sales" else 7).items():
            out.setdefault(k, []).extend(v)
    return out


def w_dept_summary(b, a, p):
    want = {}
    for r in data(b):
        want[REPS[r[2]]] = want.get(REPS[r[2]], 0) + rev(r)
    got = all_pairs(p)
    ok = all(any(near(v, want[d]) for v in got.get(d, [])) for d in want)
    return ok, f"want {want} got {{k: got.get(k) for k in want}}"


def w_commission(b, a, p):
    j = find_col(a, "commission")
    if j is None:
        return False, "no Commission column"
    d = data(a)
    bad = [r for r in d if not near(r[j], rev(r) * DEPT_RATE.get(REPS[r[2]], 0.03))]
    tot = next(r for r in a if r and r[0] == "Total")
    ok = not bad and near(tot[j], sum(rev(r) * DEPT_RATE.get(REPS[r[2]], 0.03) for r in d))
    return ok, f"{len(bad)} wrong rows, total {tot[j] if j < len(tot) else None}"


def w_above(b, a, p):
    j = find_col(a, "above") or find_col(a, "vs") or find_col(a, "average")
    if j is None:
        j = next((k for k in range(7, len(a[1])) if a[1][k] in ("Above", "Below")), None)
    if j is None:
        return False, "no Above/Below column"
    d = data(a)
    avg = {g: sum(rev(r) for r in d if r[1] == g) / max(1, sum(1 for r in d if r[1] == g)) for g in REGIONS}
    bad = [r for r in d if str(r[j]).strip().lower() != ("above" if rev(r) > avg[r[1]] else "below")]
    return not bad, f"{len(bad)} wrong"


def w_region_share(b, a, p):
    j = find_col(a, "share")
    if j is None:
        return False, "no share column"
    d = data(a)
    tot = {g: sum(rev(r) for r in d if r[1] == g) for g in REGIONS}
    bad = [r for r in d if not near(r[j], rev(r) / tot[r[1]] if tot[r[1]] else 0)]
    return not bad, f"{len(bad)} wrong"


def w_quarter(b, a, p):
    j = find_col(a, "quarter")
    if j is None:
        return False, "no Quarter column"
    d = data(a)
    q = lambda r: f"Q{(as_date(r[0]).month - 1) // 3 + 1}"
    bad = [r for r in d if str(r[j]).strip() != q(r)]
    want = {}
    for r in d:
        want[q(r)] = want.get(q(r), 0) + rev(r)
    got = pairs(a, j + 1)
    ok = not bad and all(any(near(v, want[k]) for v in got.get(k, [])) for k in want)
    return ok, f"{len(bad)} wrong quarters; want {want}"


def w_missing(b, a, p):
    j = find_col(a, "status") or find_col(a, "missing")
    if j is None:
        return False, "no status column"
    d = data(a)
    bad = [r for r in d if str(r[j]).strip().lower() != ("missing" if r[4] is None else "ok")]
    n = sum(1 for r in data(b) if r[4] is None)
    return not bad and near(cell(a, "J10"), n), f"{len(bad)} wrong, J10={cell(a, 'J10')} want {n}"


def w_top_rep(b, a, p):
    tot = {}
    for r in data(b):
        tot[r[2]] = tot.get(r[2], 0) + rev(r)
    best = max(tot, key=tot.get)
    return cell(a, "J12") == best, f"J12={cell(a, 'J12')} want {best}"


def w_hindi_gst(b, a, p):
    for j in range(6, len(a[1])):  # a new column, or the existing "Revenue with GST" column put right
        d = data(a)
        if d and all(near(r[j], rev(r) * 1.18) for r in d if j < len(r)):
            return True, f"column {sheets.col_name(j + 1)}"
    return False, "no column equal to revenue × 1.18"


def w_band(b, a, p):
    j = find_col(a, "band")
    if j is None:
        return False, "no Band column"
    band = lambda u: "XL" if u >= 150 else "Large" if u >= 100 else "Medium" if u >= 50 else "Small"
    bad = [r for r in data(a) if r[j] != band(r[3])]
    return not bad, f"{len(bad)} wrong"


def w_workdays(b, a, p):
    j = find_col(a, "working") or find_col(a, "days")
    if j is None:
        return False, "no working days column"
    end = dt.date(2025, 6, 30)
    def nwd(s):
        n, d = 0, s
        while d <= end:
            n += d.weekday() < 5
            d += dt.timedelta(days=1)
        return n
    bad = [r for r in data(a) if not near(r[j], nwd(as_date(r[0])))]
    return not bad, f"{len(bad)} wrong"


WORKFLOWS = [
    ("Look up each rep's department from the Staff sheet, then make a summary of total revenue by department", w_dept_summary),
    ("Add a Commission column: 5% of revenue for reps in the Sales department, 8% for Enterprise and 3% for everyone else "
     "(take each rep's department from the Staff sheet), and put the commission total in the totals row", w_commission),
    ("Add a column that says Above if a sale's revenue is above the average revenue of its region, otherwise Below", w_above),
    ("Add a column with each sale's share of its region's total revenue", w_region_share),
    ("Add a Quarter column (Q1, Q2, ...) from the sale date, then make a summary of total revenue by quarter", w_quarter),
    ("Add a Price Status column that says Missing when the unit price is blank and OK otherwise, and in J10 count how many "
     "prices are missing", w_missing),
    ("Find the rep with the highest total revenue and put their name in J12", w_top_rep),
    ("हर बिक्री के लिए 18% GST जोड़कर कुल राशि का एक नया कॉलम बनाओ", w_hindi_gst),
    ("Add a Band column: look up each sale's units in the Bands sheet with an approximate match (the band whose Min Units "
     "is the largest one not above the units)", w_band),
    ("Add a column with the number of working days from the sale date to 30 June 2025 (both days counted)", w_workdays),
]

TASKS = [
    ("Add a column showing each sale's share of total revenue", t_share),
    ("In cell J2 put the total revenue for the North region in February 2025", t_sumifs),
    ("Add a Department column for each rep, looked up from the Staff sheet", t_lookup),
    ("Add a running total of revenue", t_running),
    ("Add a Rank column ranking each sale by revenue, 1 = highest", t_rank),
    ("Add a Size column: Large if units are 100 or more, Medium if 50 to 99, otherwise Small", t_band),
    ("Make a summary of total revenue and number of sales by region", t_pivot),
    ("Add a Due Date column: the last day of the month after the sale date", t_due),
    ("Add a First Name column taken from the Rep name", t_first),
    ("In J5 calculate the monthly payment on a loan of 500000 at 9% a year over 60 months", t_pmt),
    ("In J8 put the average unit price, ignoring blank prices", t_avg_blank),
    ("Column G shows #REF! errors. Fix it: it should be revenue plus 18% GST", t_fix),
    ("Add a sale: 2025-03-15, North, Asha Verma, 40 units at 250", t_add_row),
    ("Sort the sales by revenue, largest first", t_sort),
    ("In the Years sheet, put the CAGR from 2019 to 2024 in D2", t_cagr),
]


def run_one(base: Path, instruction: str):
    tmp = Path(tempfile.mkdtemp())
    p = tmp / "sales.xlsx"
    shutil.copy2(base, p)
    from engine.office import recalc
    before = values(p)
    gen = sheets.edit(p, instruction)
    try:
        while True:
            next(gen)
    except StopIteration as stop:
        summary, notes = stop.value
    return p, before, summary, notes


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", type=int, default=None)
    ap.add_argument("--out", default="evals/results_sheets.json")
    ap.add_argument("--set", default="tasks", choices=["tasks", "workflows", "all"])
    a = ap.parse_args()
    chosen = TASKS if a.set == "tasks" else WORKFLOWS if a.set == "workflows" else TASKS + WORKFLOWS
    base = Path(tempfile.mkdtemp()) / "base.xlsx"
    build(base)
    from engine.office import recalc
    recalc(base, {"Sales"})
    results = []
    for i, (instruction, check) in enumerate(chosen, 1):
        if a.only and i != a.only:
            continue
        t0 = time.time()
        try:
            p, before, summary, notes = run_one(base, instruction)
            ok, note = check(before, values(p), p)
            err = None
        except (sheets.StepError, OfficeError, Exception) as e:
            ok, note, notes, err = False, "", [], f"{type(e).__name__}: {e}"
        secs = round(time.time() - t0, 1)
        results.append({"task": instruction, "ok": ok, "note": note, "changes": notes, "error": err, "seconds": secs})
        print(f"{'PASS' if ok else 'FAIL'} {secs:5}s  {instruction}\n      {err or note}", flush=True)
    passed = sum(r["ok"] for r in results)
    print(f"\n{passed}/{len(results)} passed")
    Path(a.out).write_text(json.dumps({"passed": passed, "total": len(results), "results": results}, indent=1, ensure_ascii=False))


if __name__ == "__main__":
    main()
